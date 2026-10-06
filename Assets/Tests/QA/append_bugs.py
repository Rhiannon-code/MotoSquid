#!/usr/bin/env python3
"""Append QA findings from Unity to MotoSquid_BugTracker.xlsx."""

import sys
import json
import datetime
import openpyxl

SHEET_NAME = "Bug Tracker"
# Column indices (1-based, matching Bug Tracker sheet)
COL_ID          = 1
COL_TITLE       = 2
COL_CATEGORY    = 3
COL_PRIORITY    = 4
COL_STATUS      = 5
COL_ASSIGNED    = 6
COL_SCENE       = 7
COL_STEPS       = 8
COL_EXPECTED    = 9
COL_ACTUAL      = 10
COL_DATE_REP    = 11
COL_DATE_RES    = 12
COL_NOTES       = 13
COL_COMMIT      = 14


def get_next_bug_num(ws):
    max_num = 0
    for row in ws.iter_rows(min_row=3, values_only=True):
        bug_id = row[COL_ID - 1]
        if bug_id and str(bug_id).startswith("BUG-"):
            try:
                num = int(str(bug_id).replace("BUG-", ""))
                max_num = max(max_num, num)
            except ValueError:
                pass
    return max_num + 1


def is_duplicate(ws, title, scene_component):
    for row in ws.iter_rows(min_row=3, values_only=True):
        existing_title  = str(row[COL_TITLE - 1])  if row[COL_TITLE - 1]  else ""
        existing_status = str(row[COL_STATUS - 1]) if row[COL_STATUS - 1] else ""
        existing_scene  = str(row[COL_SCENE - 1])  if row[COL_SCENE - 1]  else ""
        if (existing_title == title
                and existing_scene == scene_component
                and existing_status not in ("Fixed", "Closed", "Won't Fix")):
            return True
    return False


def main():
    if len(sys.argv) < 3:
        print("Usage: append_bugs.py <findings.json> <bugtracker.xlsx>", file=sys.stderr)
        sys.exit(1)

    findings_path = sys.argv[1]
    tracker_path  = sys.argv[2]

    with open(findings_path, encoding="utf-8-sig") as f:
        findings = json.load(f)

    wb = openpyxl.load_workbook(tracker_path)
    if SHEET_NAME not in wb.sheetnames:
        print(f"Sheet '{SHEET_NAME}' not found in workbook.", file=sys.stderr)
        sys.exit(1)

    ws = wb[SHEET_NAME]
    today      = datetime.date.today()
    next_num   = get_next_bug_num(ws)
    added      = 0
    skipped    = 0

    for finding in findings:
        title          = finding.get("title", "")
        scene_comp     = finding.get("sceneComponent", "")

        if is_duplicate(ws, title, scene_comp):
            skipped += 1
            continue

        bug_id = f"BUG-{next_num:03d}"
        next_num += 1

        row_data = [""] * 14
        row_data[COL_ID       - 1] = bug_id
        row_data[COL_TITLE    - 1] = title
        row_data[COL_CATEGORY - 1] = finding.get("category", "")
        row_data[COL_PRIORITY - 1] = finding.get("priority", "")
        row_data[COL_STATUS   - 1] = "Open"
        row_data[COL_ASSIGNED - 1] = finding.get("assignedTo", "")
        row_data[COL_SCENE    - 1] = scene_comp
        row_data[COL_STEPS    - 1] = finding.get("stepsToReproduce", "")
        row_data[COL_EXPECTED - 1] = finding.get("expectedBehaviour", "")
        row_data[COL_ACTUAL   - 1] = finding.get("actualBehaviour", "")
        row_data[COL_DATE_REP - 1] = today
        row_data[COL_DATE_RES - 1] = None
        row_data[COL_NOTES    - 1] = finding.get("notes", "")
        row_data[COL_COMMIT   - 1] = None

        ws.append(row_data)
        added += 1

    wb.save(tracker_path)

    msg = f"Added {added} bug(s)"
    if skipped:
        msg += f", skipped {skipped} duplicate(s) already tracked as Open"
    print(msg)


if __name__ == "__main__":
    main()
